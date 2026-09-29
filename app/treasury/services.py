"""Treasury/bank services. Every movement posts a journal entry (spec §6, §1.1).

Balances are read from the ledger (posted journal lines on the entity's GL
account), never stored — so there is only ever one source of truth.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

from flask import current_app

from app.accounting import posting
from app.accounting.models import Account, JournalEntry, JournalLine
from app.core import audit, fx, settings
from app.core.money import quantize_amount, quantize_rate, to_decimal
from app.core.numbering import next_number
from app.extensions import db
from app.treasury.models import BankAccount, Treasury, TreasuryTxn


class CurrencyLockedError(Exception):
    pass


class BalanceNotZeroError(Exception):
    pass


class ApprovalPendingError(Exception):
    pass


# --- entity creation ------------------------------------------------------
def _cash_parent_id():
    acc = db.session.scalar(db.select(Account).filter_by(code="11"))
    return acc.id if acc else None


def create_treasury(*, name_ar, branch_id, currency_code, type="main",
                    opening_balance=0, user_id=None) -> Treasury:
    t = Treasury(
        name_ar=name_ar, branch_id=branch_id, currency_code=currency_code,
        type=type, opening_balance=quantize_amount(opening_balance),
        created_by_id=user_id,
    )
    t.account = Account(
        code=f"115{_next_seq('115')}", name_ar=f"خزينة: {name_ar}",
        name_en=name_ar, type="asset", is_postable=True,
        parent_id=_cash_parent_id(), created_by_id=user_id,
    )
    db.session.add(t)
    db.session.flush()
    _post_opening(t, user_id)
    audit.record(action="treasury.create", entity="treasury", entity_id=t.id,
                 new={"name": name_ar, "currency": currency_code})
    return t


def create_bank_account(*, name_ar, bank_name, account_number, branch_id,
                        currency_code, opening_balance=0, user_id=None) -> BankAccount:
    b = BankAccount(
        name_ar=name_ar, bank_name=bank_name, account_number=account_number,
        branch_id=branch_id, currency_code=currency_code,
        opening_balance=quantize_amount(opening_balance), created_by_id=user_id,
    )
    b.account = Account(
        code=f"116{_next_seq('116')}", name_ar=f"بنك: {name_ar}",
        name_en=name_ar, type="asset", is_postable=True,
        parent_id=_cash_parent_id(), created_by_id=user_id,
    )
    db.session.add(b)
    db.session.flush()
    _post_opening(b, user_id)
    audit.record(action="bank.create", entity="bank_account", entity_id=b.id,
                 new={"name": name_ar, "currency": currency_code})
    return b


def _next_seq(prefix: str) -> str:
    """A short unique numeric suffix for an auto-created GL account code."""
    count = db.session.scalar(
        db.select(db.func.count(Account.id)).filter(Account.code.like(f"{prefix}%"))
    )
    return f"{(count or 0) + 1:03d}"


def _post_opening(entity, user_id):
    """Opening balance -> opening journal entry (spec §4.4)."""
    ob = quantize_amount(entity.opening_balance)
    if ob == 0:
        return
    equity_id = posting.account_for("opening.equity")
    rate = fx.rate_to_book(entity.currency_code) or quantize_rate(1)
    e = posting.build_entry(
        entry_date=date.today(), branch_id=entity.branch_id,
        source_doc_type="opening", source_doc_id=entity.id,
        memo=f"رصيد افتتاحي — {entity.name_ar}", user_id=user_id,
        lines=[
            posting.line(entity.account_id, currency=entity.currency_code,
                         amount=ob, side="debit", fx_rate=rate),
            posting.line(equity_id, currency=entity.currency_code,
                         amount=ob, side="credit", fx_rate=rate),
        ],
    )
    posting.post_entry(e, user_id=user_id)


# --- balances (from the ledger) -------------------------------------------
def balance_native(account_id: int, up_to: date | None = None) -> Decimal:
    """Balance in the entity's own currency (spec §6.6 — no conversion)."""
    q = (
        db.select(
            db.func.coalesce(db.func.sum(JournalLine.debit_original), 0),
            db.func.coalesce(db.func.sum(JournalLine.credit_original), 0),
        )
        .join(JournalEntry, JournalEntry.id == JournalLine.entry_id)
        .filter(JournalLine.account_id == account_id)
        .filter(JournalEntry.status == "posted")
    )
    if up_to:
        q = q.filter(JournalEntry.date <= up_to)
    dr, cr = db.session.execute(q).one()
    return quantize_amount(to_decimal(dr) - to_decimal(cr))


def has_movements(entity) -> bool:
    """Any posted line on the entity's account = a movement (locks currency)."""
    n = db.session.scalar(
        db.select(db.func.count(JournalLine.id))
        .join(JournalEntry, JournalEntry.id == JournalLine.entry_id)
        .filter(JournalLine.account_id == entity.account_id)
        .filter(JournalEntry.status == "posted")
    )
    return (n or 0) > 0


# --- movements ------------------------------------------------------------
def _requires_approval(amount_book: Decimal, user) -> bool:
    try:
        limit = settings.get("treasury.transfer_approval_limit")
    except settings.SettingNotFoundError:
        return False
    if limit in (None, ""):
        return False
    if user and any(r.code == "owner" for r in user.roles):
        return False
    return to_decimal(amount_book) > to_decimal(limit)


def deposit(*, entity, counter_account_id, amount, entry_date=None, memo=None,
            user_id=None) -> TreasuryTxn:
    """Money into an entity from a GL source account."""
    return _simple_movement(
        kind="deposit", debit_account=entity.account_id,
        credit_account=counter_account_id, entity=entity, amount=amount,
        entry_date=entry_date, memo=memo, user_id=user_id,
    )


def withdraw(*, entity, counter_account_id, amount, entry_date=None, memo=None,
             user_id=None) -> TreasuryTxn:
    """Money out of an entity to a GL destination account."""
    return _simple_movement(
        kind="withdraw", debit_account=counter_account_id,
        credit_account=entity.account_id, entity=entity, amount=amount,
        entry_date=entry_date, memo=memo, user_id=user_id,
    )


def _simple_movement(*, kind, debit_account, credit_account, entity, amount,
                     entry_date, memo, user_id):
    amount = quantize_amount(amount)
    cur = entity.currency_code
    rate = fx.rate_to_book(cur) or quantize_rate(1)
    e = posting.build_entry(
        entry_date=entry_date or date.today(), branch_id=entity.branch_id,
        source_doc_type=f"treasury.{kind}", memo=memo, user_id=user_id,
        lines=[
            posting.line(debit_account, currency=cur, amount=amount,
                         side="debit", fx_rate=rate),
            posting.line(credit_account, currency=cur, amount=amount,
                         side="credit", fx_rate=rate),
        ],
    )
    posting.post_entry(e, user_id=user_id)
    txn = TreasuryTxn(
        number=next_number(entity.branch_id, "CASH", prefix="CSH-"),
        date=e.date, kind=kind, src_account_id=credit_account,
        dst_account_id=debit_account, amount_original=amount,
        currency_code=cur, fx_rate=rate, status="posted",
        journal_entry_id=e.id, memo=memo, created_by_id=user_id,
    )
    db.session.add(txn)
    db.session.flush()
    audit.record(action=f"treasury.{kind}", entity="treasury_txn",
                 entity_id=txn.id, new={"amount": str(amount), "currency": cur},
                 branch_id=entity.branch_id)
    return txn


def internal_transfer(*, src, dst, amount, entry_date=None, memo=None,
                      user_id=None, user=None) -> TreasuryTxn:
    """Same-currency transfer between two entities (spec §6). No FX."""
    if src.currency_code != dst.currency_code:
        raise ValueError("التحويل الداخلي يتطلب نفس العملة. استخدم التحويل بعملتين.")
    amount = quantize_amount(amount)
    cur = src.currency_code
    rate = fx.rate_to_book(cur) or quantize_rate(1)
    amount_book = quantize_amount(amount * rate)

    txn = TreasuryTxn(
        number=next_number(src.branch_id, "CASH", prefix="CSH-"),
        date=entry_date or date.today(), kind="internal_transfer",
        src_account_id=src.account_id, dst_account_id=dst.account_id,
        amount_original=amount, currency_code=cur, fx_rate=rate,
        memo=memo, created_by_id=user_id,
    )

    if _requires_approval(amount_book, user):
        txn.status = "pending_approval"
        db.session.add(txn)
        db.session.flush()
        audit.record(action="treasury.transfer.pending", entity="treasury_txn",
                     entity_id=txn.id, new={"amount": str(amount)})
        return txn

    _post_transfer_entry(txn, src, dst, user_id)
    return txn


def _post_transfer_entry(txn, src, dst, user_id):
    cur = src.currency_code
    e = posting.build_entry(
        entry_date=txn.date, branch_id=src.branch_id,
        source_doc_type="treasury.internal_transfer", memo=txn.memo,
        user_id=user_id,
        lines=[
            posting.line(dst.account_id, currency=cur, amount=txn.amount_original,
                         side="debit", fx_rate=txn.fx_rate),
            posting.line(src.account_id, currency=cur, amount=txn.amount_original,
                         side="credit", fx_rate=txn.fx_rate),
        ],
    )
    posting.post_entry(e, user_id=user_id)
    txn.status = "posted"
    txn.journal_entry_id = e.id
    db.session.add(txn)
    db.session.flush()
    audit.record(action="treasury.internal_transfer", entity="treasury_txn",
                 entity_id=txn.id, new={"amount": str(txn.amount_original)})


def fx_transfer(*, src, dst, sent_amount, received_amount, src_rate=None,
                dst_rate=None, bank_fee=0, entry_date=None, memo=None,
                user_id=None, user=None) -> TreasuryTxn:
    """Cross-currency transfer (spec §6 example: UAE -> Egypt).

    Books the receiving side at its own rate, the sending side at its rate, and
    plugs the book-currency difference to FX gain/loss (spec §13). An optional
    bank fee is booked as an expense.
    """
    sent = quantize_amount(sent_amount)
    recv = quantize_amount(received_amount)
    fee = quantize_amount(bank_fee or 0)
    r_src = fx.require_rate(src.currency_code, current_app.config["BOOK_CURRENCY"],
                            manual=src_rate)
    r_dst = fx.require_rate(dst.currency_code, current_app.config["BOOK_CURRENCY"],
                            manual=dst_rate)
    book_sent = quantize_amount(sent * r_src)
    book_recv = quantize_amount(recv * r_dst)
    book_fee = fee  # fee entered in book currency

    txn = TreasuryTxn(
        number=next_number(src.branch_id, "CASH", prefix="CSH-"),
        date=entry_date or date.today(), kind="fx_transfer",
        src_account_id=src.account_id, dst_account_id=dst.account_id,
        amount_original=sent, currency_code=src.currency_code, fx_rate=r_src,
        dst_currency_code=dst.currency_code, dst_fx_rate=r_dst,
        received_amount=recv, bank_fee=fee, memo=memo, created_by_id=user_id,
    )

    if _requires_approval(book_sent, user):
        txn.status = "pending_approval"
        db.session.add(txn)
        db.session.flush()
        audit.record(action="treasury.transfer.pending", entity="treasury_txn",
                     entity_id=txn.id, new={"sent": str(sent)})
        return txn

    _post_fx_entry(txn, src, dst, book_sent, book_recv, book_fee, user_id)
    return txn


def _post_fx_entry(txn, src, dst, book_sent, book_recv, book_fee, user_id):
    lines = [
        posting.line(dst.account_id, currency=dst.currency_code,
                     amount=txn.received_amount, side="debit", fx_rate=txn.dst_fx_rate),
        posting.line(src.account_id, currency=src.currency_code,
                     amount=txn.amount_original, side="credit", fx_rate=txn.fx_rate),
    ]
    if book_fee and book_fee > 0:
        lines.append(posting.line(
            posting.account_for("bank.fee"),
            currency=current_app.config["BOOK_CURRENCY"], amount=book_fee,
            side="debit", fx_rate=1))

    # Plug the book difference to FX gain/loss so the entry balances (§13).
    debit_total = book_recv + (book_fee or 0)
    credit_total = book_sent
    diff = quantize_amount(credit_total - debit_total)
    if diff > 0:
        lines.append(posting.line(
            posting.account_for("fx.loss"),
            currency=current_app.config["BOOK_CURRENCY"], amount=diff,
            side="debit", fx_rate=1))
    elif diff < 0:
        lines.append(posting.line(
            posting.account_for("fx.gain"),
            currency=current_app.config["BOOK_CURRENCY"], amount=-diff,
            side="credit", fx_rate=1))

    e = posting.build_entry(
        entry_date=txn.date, branch_id=src.branch_id,
        source_doc_type="treasury.fx_transfer", memo=txn.memo, user_id=user_id,
        lines=lines,
    )
    posting.post_entry(e, user_id=user_id)
    txn.status = "posted"
    txn.journal_entry_id = e.id
    db.session.add(txn)
    db.session.flush()
    audit.record(action="treasury.fx_transfer", entity="treasury_txn",
                 entity_id=txn.id,
                 new={"sent": str(txn.amount_original), "recv": str(txn.received_amount)})


def approve_and_post(txn: TreasuryTxn, *, user_id):
    """Owner approves a pending transfer, which then posts (spec §6 limits)."""
    if txn.status != "pending_approval":
        return txn
    src = _entity_by_account(txn.src_account_id)
    dst = _entity_by_account(txn.dst_account_id)
    txn.approved_by_id = user_id
    if txn.kind == "fx_transfer":
        book_sent = quantize_amount(txn.amount_original * txn.fx_rate)
        book_recv = quantize_amount(txn.received_amount * txn.dst_fx_rate)
        _post_fx_entry(txn, src, dst, book_sent, book_recv,
                       quantize_amount(txn.bank_fee or 0), user_id)
    else:
        _post_transfer_entry(txn, src, dst, user_id)
    audit.record(action="treasury.transfer.approve", entity="treasury_txn",
                 entity_id=txn.id, new={"approved_by": user_id})
    return txn


def _entity_by_account(account_id):
    t = db.session.scalar(db.select(Treasury).filter_by(account_id=account_id))
    if t:
        return t
    return db.session.scalar(db.select(BankAccount).filter_by(account_id=account_id))


def deactivate(entity, *, user_id):
    """Suspend a treasury/bank — only if its balance is zero (spec §16.17)."""
    if balance_native(entity.account_id) != 0:
        raise BalanceNotZeroError("لا يمكن تعطيل خزينة/حساب عليه رصيد. يجب أن يكون صفرًا.")
    entity.is_active = False
    audit.record(action=f"{entity.kind}.deactivate",
                 entity=entity.kind, entity_id=entity.id)


def statement(account_id, *, up_to=None):
    """Ledger statement for an entity's account: each posted line + running
    native balance (spec §6 reports)."""
    q = (
        db.select(JournalLine, JournalEntry)
        .join(JournalEntry, JournalEntry.id == JournalLine.entry_id)
        .filter(JournalLine.account_id == account_id)
        .filter(JournalEntry.status == "posted")
        .order_by(JournalEntry.date, JournalEntry.id)
    )
    if up_to:
        q = q.filter(JournalEntry.date <= up_to)
    rows, running = [], Decimal("0")
    for line, entry in db.session.execute(q).all():
        delta = quantize_amount(to_decimal(line.debit_original)
                                - to_decimal(line.credit_original))
        running = quantize_amount(running + delta)
        rows.append({"entry": entry, "line": line, "delta": delta, "balance": running})
    return rows


def import_statement(bank, file_storage, *, user_id=None):
    """Parse a bank statement (xlsx) and auto-match rows to unreconciled ledger
    lines by amount (spec TRS-07). Returns (matched, unmatched)."""
    from decimal import Decimal
    from openpyxl import load_workbook
    from app.treasury.models import ReconItem

    wb = load_workbook(file_storage, data_only=True)
    ws = wb.active
    # expect columns: date, description, amount (any order-tolerant: find a number)
    stmt = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        amount = None
        for v in row:
            try:
                amount = to_decimal(v)
                break
            except Exception:
                continue
        if amount is not None:
            stmt.append(amount)

    rows = statement(bank.account_id)  # ledger lines with deltas
    cleared = {r.journal_line_id for r in db.session.scalars(
        db.select(ReconItem).filter_by(bank_account_id=bank.id))}
    matched = 0
    for amt in stmt:
        for r in rows:
            if r["line"].id in cleared:
                continue
            if quantize_amount(r["delta"]) == quantize_amount(amt):
                db.session.add(ReconItem(bank_account_id=bank.id,
                                         journal_line_id=r["line"].id,
                                         reconciled_by_id=user_id,
                                         statement_ref="import"))
                cleared.add(r["line"].id)
                matched += 1
                break
    db.session.flush()
    return matched, len(stmt) - matched
