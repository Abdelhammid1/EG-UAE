"""M13 — Excel opening-balance migration (MIG-02..04)."""
import io
import json
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app import create_app
from app.cli import run_seed
from app.extensions import db
from config import TestConfig


@pytest.fixture
def app():
    app = create_app(TestConfig)
    with app.app_context():
        run_seed(create_tables=True)
        yield app
        db.session.remove()
        db.drop_all()


def _uid():
    from app.auth.models import User
    return db.session.scalar(db.select(User).filter_by(username="owner")).id


def _xlsx(headers, rows):
    wb = Workbook(); ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def _batch(kind, rows):
    from app.migration import services
    from app.migration.models import ImportBatch
    rows, errs = services.validate(kind, rows)
    b = ImportBatch(kind=kind, rows_json=json.dumps(rows, default=str),
                    row_count=len(rows), error_count=errs, status="validated")
    db.session.add(b); db.session.flush()
    return b, errs


# --- templates ------------------------------------------------------------
def test_template_has_headers_and_lists(app):
    from app.migration import services
    buf = services.generate_template("products")
    wb = load_workbook(buf)
    assert "القيم المسموحة" in wb.sheetnames
    ws = wb[wb.sheetnames[0]]
    assert ws.cell(row=1, column=1).value == "اسم المنتج"


# --- validation -----------------------------------------------------------
def test_validate_flags_bad_currency(app):
    from app.migration import services
    rows = services.parse_file("customers", _xlsx(
        ["الاسم", "الهاتف", "النوع", "العملة", "حد", "رصيد"],
        [["عميل", "050", "retail", "XXX", "0", "100"]]))
    rows, errs = services.validate("customers", rows)
    assert errs == 1
    assert any("عملة" in e for e in rows[0]["_errors"])


def test_validate_flags_duplicate_barcode(app):
    from app.migration import services
    rows = services.parse_file("products", _xlsx(
        ["a", "b", "c", "d", "e", "f", "g"],
        [["منتج1", "9001", "", "قطعة", "100", "50", "FIFO"],
         ["منتج2", "9001", "", "قطعة", "100", "50", "FIFO"]]))
    rows, errs = services.validate("products", rows)
    assert errs == 1  # second row duplicate barcode


# --- commit ---------------------------------------------------------------
def test_commit_creates_customers_with_opening_balance(app):
    from app.migration import services
    from app.sales.models import Customer
    from app.sales import services as sales
    from app.accounting.services import trial_balance
    rows = services.parse_file("customers", _xlsx(
        ["الاسم", "الهاتف", "النوع", "العملة", "حد", "رصيد"],
        [["عميل افتتاحي", "050", "wholesale", "AED", "10000", "3000"]]))
    b, errs = _batch("customers", rows)
    assert errs == 0
    services.commit(b, user_id=_uid())
    db.session.commit()
    c = db.session.scalar(db.select(Customer).filter_by(name_ar="عميل افتتاحي"))
    assert c is not None
    assert sales.ar_balance(c) == Decimal("3000.0000")
    assert trial_balance()["balanced"]


def test_commit_creates_products_then_stock(app):
    from app.migration import services
    from app.inventory.models import Product, Warehouse
    from app.inventory import services as inv
    from app.core.models import Branch
    # a warehouse to receive into
    b = db.session.scalar(db.select(Branch))
    wh = Warehouse(name_ar="مخزن الافتتاح", branch_id=b.id)
    db.session.add(wh); db.session.commit()
    # products
    prows = services.parse_file("products", _xlsx(
        ["a", "b", "c", "d", "e", "f", "g"],
        [["شاشة", "8001", "", "قطعة", "500", "400", "FIFO"]]))
    pb, e1 = _batch("products", prows); assert e1 == 0
    services.commit(pb, user_id=_uid()); db.session.commit()
    # stock
    srows = services.parse_file("stock", _xlsx(
        ["a", "b", "c", "d", "e"],
        [["8001", "مخزن الافتتاح", "5", "450", ""]]))
    sb, e2 = _batch("stock", srows); assert e2 == 0
    services.commit(sb, user_id=_uid()); db.session.commit()
    p = db.session.scalar(db.select(Product).filter_by(barcode="8001"))
    assert inv.qty_on_hand(p.id, wh.id) == Decimal("5.0000")


def test_commit_is_all_or_nothing(app):
    """A failure mid-batch rolls the whole thing back (MIG-03.5)."""
    from app.migration import services
    from app.migration.models import ImportBatch
    from app.inventory.models import Product, Warehouse
    from app.inventory import services as inv
    from app.core.models import Branch
    br = db.session.scalar(db.select(Branch))
    wh = Warehouse(name_ar="مخزن", branch_id=br.id)
    p = Product(name_ar="سلعة", barcode="7001", valuation_method="FIFO")
    db.session.add_all([wh, p]); db.session.commit()
    # row 1 valid, row 2 references a product that doesn't exist -> raises at commit
    rows = [
        {"barcode": "7001", "warehouse": "مخزن", "qty": "5", "unit_cost": "100",
         "serials": "", "_line": 2},
        {"barcode": "NOPE", "warehouse": "مخزن", "qty": "3", "unit_cost": "100",
         "serials": "", "_line": 3},
    ]
    b = ImportBatch(kind="stock", rows_json=json.dumps(rows), row_count=2,
                    error_count=0, status="validated")
    db.session.add(b); db.session.flush()
    with pytest.raises(Exception):
        services.commit(b, user_id=_uid())
    db.session.rollback()
    # the first row's stock was NOT persisted (whole batch rolled back)
    assert inv.qty_on_hand(p.id, wh.id) == Decimal("0.0000")
